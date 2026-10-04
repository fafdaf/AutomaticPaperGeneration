# -*- coding: utf-8 -*-
"""多来源要求：返回 (文本列表, 图片路径列表)"""
import os


def read_text_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".md"):
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        try:
            import pdfplumber
            out = []
            with pdfplumber.open(path) as pdf:
                for pg in pdf.pages:
                    t = pg.extract_text()
                    if t:
                        out.append(t)
            if out:
                return "\n".join(out)
        except ImportError:
            pass
        from PyPDF2 import PdfReader
        reader = PdfReader(path)
        return "\n".join((pg.extract_text() or "") for pg in reader.pages)
    raise ValueError(f"不支持的文本格式：{ext}")


def classify_paths(paths):
    text_blocks, images = [], []
    for p in paths:
        if not os.path.exists(p):
            print(f"[警告] 文件不存在，已跳过：{p}")
            continue
        ext = os.path.splitext(p)[1].lower()
        if ext in (".png", ".jpg", ".jpeg", ".bmp", ".webp", ".gif"):
            images.append(p)
        elif ext == ".doc":
            print(f"[警告] 不支持旧版 .doc 格式，请另存为 .docx：{p}")
        else:
            try:
                text_blocks.append(read_text_file(p))
            except Exception as e:
                print(f"[错误] 解析 {p} 失败：{e}")
    return text_blocks, images